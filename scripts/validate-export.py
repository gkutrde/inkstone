from pathlib import Path
from docx import Document
from lxml import html as lxml_html
import json
import zipfile

ROOT = Path(__file__).resolve().parent.parent
QA = ROOT.parent / "analysis" / "export-qa"
docx_path = QA / "Inkstone-export.docx"
html_path = QA / "Inkstone-export.html"
markdown_copy_path = QA / "Inkstone-copy.md"
markdown_path = QA / "sample.md"
report_path = QA / "report.json"
report = json.loads(report_path.read_text(encoding="utf-8"))

assert docx_path.exists() and docx_path.stat().st_size > 1000
with zipfile.ZipFile(docx_path) as archive:
    assert archive.testzip() is None
    names = archive.namelist()
    styles_xml = archive.read("word/styles.xml").decode("utf-8")
    numbering_xml = archive.read("word/numbering.xml").decode("utf-8")
    assert "Microsoft YaHei" in styles_xml
    assert "https://example.com" in archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert any(name.startswith("word/media/") for name in names)
    assert "w:start w:val=\"3\"" in numbering_xml
    assert "w:gridCol w:w=\"4800\"" in archive.read("word/document.xml").decode("utf-8")

document = Document(docx_path)
paragraph_text = "\n".join(p.text for p in document.paragraphs)
assert "导出验收" in paragraph_text
assert "第一段包含" in paragraph_text
assert "待处理事项" in paragraph_text
assert "Hello, 世界" in paragraph_text
assert any(p.style.name.startswith("Heading") or p.style.name == "Title" for p in document.paragraphs)
assert any(p.style.name == "Code Block" for p in document.paragraphs)
assert any(p._p.pPr is not None and p._p.pPr.numPr is not None for p in document.paragraphs)
first_paragraph = next(p for p in document.paragraphs if "第一段包含" in p.text)
assert any(run.bold for run in first_paragraph.runs)
assert any(run.italic for run in first_paragraph.runs)
assert len(document.tables) == 1
assert len(document.tables[0].rows) == 3
assert document.tables[0].cell(1, 0).text == "中文内容"
assert len(document.inline_shapes) == 1

html_document = lxml_html.parse(str(html_path)).getroot()
assert "导出验收" in html_document.text_content()
assert len(html_document.xpath("//table")) == 1
assert len(html_document.xpath("//img")) == 1
assert markdown_copy_path.read_text(encoding="utf-8") == markdown_path.read_text(encoding="utf-8")

report.update({
    "docxBytes": docx_path.stat().st_size,
    "paragraphs": len(document.paragraphs),
    "tables": len(document.tables),
    "images": len(document.inline_shapes),
    "structureChecks": "passed",
})
report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
